from pptx import Presentation
prs = Presentation('test_table_layout.pptx')
for si, sl in enumerate(prs.slides):
    print(f'\n=== Slide {si+1} ===')
    for sh in sl.shapes:
        left_in = round(sh.left / 914400, 2)
        top_in  = round(sh.top  / 914400, 2)
        w_in    = round(sh.width / 914400, 2)
        h_in    = round(sh.height / 914400, 2)
        is_tbl  = hasattr(sh, 'table')
        mark    = ' [TABLE]' if is_tbl else ''
        print(f'  {sh.name:30s}{mark} top={top_in}", left={left_in}", w={w_in}", h={h_in}"')
