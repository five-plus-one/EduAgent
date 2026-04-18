import os
import io

try:
    from pypdf import PdfReader
except ImportError:
    pass

try:
    import docx
except ImportError:
    pass

try:
    from pptx import Presentation as PptxPresentation
except ImportError:
    PptxPresentation = None


def parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file.

    Raises:
        ValueError: encrypted PDF, scanned-image PDF (no extractable text), or corrupted file.
    """
    try:
        reader = PdfReader(file_path)
    except Exception as e:
        err = str(e).lower()
        if "encrypt" in err or "password" in err:
            raise ValueError(
                "PDF 文件已加密（有密码保护），请先移除密码后重新上传。"
            )
        raise ValueError(f"无法读取 PDF 文件（文件可能已损坏）：{e}")

    if reader.is_encrypted:
        raise ValueError(
            "PDF 文件已加密（有密码保护），请先移除密码后重新上传。"
        )

    text = ""
    for page in reader.pages:
        try:
            t = page.extract_text()
            if t:
                text += t + "\n"
        except Exception:
            pass  # 跳过无法提取的单页，继续处理其他页

    if not text.strip():
        raise ValueError(
            "PDF 不包含可提取的文字（可能是扫描版图片 PDF）。"
            "请上传含有可复制文字的 PDF 文件，或使用 OCR 工具处理后重新上传。"
        )

    return text


def parse_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])


def parse_txt(file_path: str) -> str:
    """Read plain-text formats: .txt / .md / .json / .csv"""
    for encoding in ("utf-8", "utf-8-sig", "gbk", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法以任何已知编码读取文件: {file_path}")


def parse_pptx(file_path: str) -> str:
    """Extract all text from a PPTX file using python-pptx."""
    if PptxPresentation is None:
        raise ValueError("python-pptx 未安装，无法解析 PPTX 文件")
    try:
        prs = PptxPresentation(file_path)
        texts = []
        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                texts.append(f"[幻灯片 {slide_idx}]\n" + "\n".join(slide_texts))
        return "\n\n".join(texts)
    except Exception as e:
        raise ValueError(f"无法解析 PPTX 文件，文件可能已损坏: {str(e)}")


def extract_text_from_file(file_path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return parse_pdf(file_path)

    elif ext in (".docx", ".doc"):
        return parse_docx(file_path)

    elif ext in (".txt", ".md", ".json", ".csv"):
        return parse_txt(file_path)

    elif ext == ".pptx":
        return parse_pptx(file_path)

    elif ext == ".ppt":
        raise ValueError(
            "安全红线拦截：系统不支持 Office 2003（.ppt）二进制格式。"
            "请在 Office 中将文件另存为 .pptx 格式后重新上传。"
        )

    else:
        raise ValueError(
            f"不支持的文件格式：{ext}。"
            "目前支持：PDF / DOCX / PPTX / TXT / MD / JSON / CSV"
        )
