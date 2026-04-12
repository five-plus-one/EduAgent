import sys
sys.path.insert(0, '.')

from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from app.services.ppt_exporter import (
    render_minimal_list, render_default, render_two_column,
    draw_bg_decor, SLIDE_W, SLIDE_H, hex2rgb,
    LAYOUT_RENDERERS, CONTENT_T, MARGIN_LEFT, CONTENT_W
)

prs = Presentation()
prs.slide_width  = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
blank = prs.slide_layouts[6]

colors = {
    'bg':  RGBColor(0x0D, 0x17, 0x2B),
    'pri': RGBColor(0x38, 0xBD, 0xF8),
    'sec': RGBColor(0x47, 0x55, 0x69),
    'acc': RGBColor(0xF5, 0x9E, 0x0B),
    'txt': RGBColor(0xF1, 0xF5, 0xF9),
}

print(f"SLIDE: {SLIDE_W:.2f} x {SLIDE_H:.2f}")
print(f"CONTENT_T={CONTENT_T:.3f}, MARGIN_LEFT={MARGIN_LEFT:.3f}, CONTENT_W={CONTENT_W:.3f}")
avail_h = SLIDE_H - CONTENT_T - 0.3
print(f"avail_h={avail_h:.3f}, 55%={avail_h*0.55:.3f}")

pages = [
    # Test 1: ONLY table
    {
        'page_index': 1, 'layout_type': 'minimal_list',
        'title': '仅含表格', 'speaker_notes': '',
        'elements': [{
            'element_id': 't1', 'type': 'table', 'position': 'full',
            'headers': ['物理量', '符号', '公式', '说明'],
            'rows': [
                ['转动惯量', 'I', r'$\frac{1}{2}mR^2$', '均质圆柱'],
                ['角动量',   'L', r'$L=I\omega$',        '刚体旋转'],
                ['力矩',     'M', r'$M=I\alpha$',        '转动第二定律'],
            ],
            'content': [], 'is_accent': False,
        }]
    },
    # Test 2: body list + table
    {
        'page_index': 2, 'layout_type': 'minimal_list',
        'title': '正文列表 + 表格', 'speaker_notes': '',
        'elements': [
            {'element_id': 'e1', 'type': 'list', 'position': 'left', 'is_accent': False,
             'content': ['转动力学研究刚体旋转运动', '类比平动：力→力矩，质量→转动惯量', '角动量守恒定律广泛应用']},
            {'element_id': 't1', 'type': 'table', 'position': 'full',
             'headers': ['物理量', '公式'],
             'rows': [['转动惯量', r'$I=\frac{1}{2}mR^2$'], ['角动量', r'$L=I\omega$']],
             'content': [], 'is_accent': False},
        ]
    },
    # Test 3: two_column + table
    {
        'page_index': 3, 'layout_type': 'two_column',
        'title': '双列布局 + 表格', 'speaker_notes': '',
        'elements': [
            {'element_id': 'e1', 'type': 'list', 'position': 'left', 'is_accent': False,
             'content': ['平动动能 E_k = mv²/2', '角速度 ω = Δθ/Δt']},
            {'element_id': 'e2', 'type': 'text_block', 'position': 'right', 'is_accent': False,
             'content': ['刚体绕固定轴转动时，转动惯量 I 由质量分布决定']},
            {'element_id': 't1', 'type': 'table', 'position': 'full',
             'headers': ['形状', '转轴', '转动惯量I'],
             'rows': [['均质圆柱', '中心轴', r'$\frac{1}{2}mR^2$'],
                      ['均质细杆', '端点',   r'$\frac{1}{3}mL^2$']],
             'content': [], 'is_accent': False},
        ]
    },
]

for page in pages:
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = colors['bg']
    draw_bg_decor(slide, page, colors)
    renderer = LAYOUT_RENDERERS.get(page['layout_type'], render_default)
    renderer(slide, page, colors)
    print(f"  Page {page['page_index']}: {page['layout_type']} | {len(slide.shapes)} shapes")

out = 'test_table_layout.pptx'
prs.save(out)
print(f"\nSaved: {out}")
